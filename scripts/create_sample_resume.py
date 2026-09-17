# from pathlib import Path
#
# import pymupdf
# from docx import Document
#
#
# OUTPUT_DIR = Path("data/raw")
# OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
#
#
# SAMPLE_TEXT = """Alex Johnson
# Software Engineer
#
# Skills:
# Python, Java, SQL, Machine Learning, Git
#
# Education:
# B.Tech in Computer Science
#
# Experience:
# Software Engineering Intern - 6 months
#
# Projects:
# Resume Screening System using Python and NLP
# """
#
#
# def create_pdf():
#     pdf_path = OUTPUT_DIR / "sample_resume.pdf"
#
#     document = pymupdf.open()
#     page = document.new_page()
#
#     page.insert_text(
#         (72, 72),
#         SAMPLE_TEXT,
#         fontsize=12,
#     )
#
#     document.save(pdf_path)
#     document.close()
#
#     return pdf_path
#
#
# def create_docx():
#     docx_path = OUTPUT_DIR / "sample_resume.docx"
#
#     document = Document()
#
#     for line in SAMPLE_TEXT.splitlines():
#         document.add_paragraph(line)
#
#     document.save(docx_path)
#
#     return docx_path
#
#
# if __name__ == "__main__":
#     pdf_path = create_pdf()
#     docx_path = create_docx()
#
#     print(f"Created PDF: {pdf_path}")
#     print(f"Created DOCX: {docx_path}")



from pathlib import Path

import pymupdf
from docx import Document
from docx.shared import Inches


OUTPUT_DIR = Path("data/raw")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


SIMPLE_RESUME = """Alex Johnson
Software Engineer

Contact
alex.johnson@example.com | +91 9876543210

Skills
Python, Java, SQL, Machine Learning, Git, Docker

Education
B.Tech in Computer Science and Engineering
ABC University | 2022 - 2026

Experience
Software Engineering Intern
Tech Solutions Pvt. Ltd. | 6 months

Worked on backend APIs and data processing using Python.

Projects
Resume Screening System
Built a resume screening system using Python and NLP.

Additional Information
Languages: English, Hindi
"""


BULLET_RESUME = """Priya Sharma
Data Science Intern

Skills
• Python
• Pandas
• NumPy
• Scikit-learn
• SQL
• Machine Learning

Experience
Data Science Intern - DataTech
• Cleaned and analyzed customer datasets.
• Built classification models using Scikit-learn.
• Created data visualizations for business reports.

Projects
• Customer Churn Prediction
• Sales Forecasting Dashboard
• Resume Matching System

Education
B.Tech - Artificial Intelligence and Data Science
XYZ University
"""


MULTI_PAGE_RESUME = """Rahul Verma
Machine Learning Engineer

Professional Summary
Machine learning enthusiast with experience in NLP,
model development, and deployment.

Skills
Python, NLP, Machine Learning, SQL, Docker, Git,
Scikit-learn, TensorFlow

Experience
Machine Learning Intern
AI Solutions
January 2025 - June 2025

Developed NLP preprocessing pipelines and evaluated
classification models.

Machine Learning Project
July 2025 - Present

Developing semantic resume-job matching using
sentence embeddings.

Education
B.Tech in Computer Science
DEF University
2022 - 2026

Certifications
Machine Learning Specialization
Cloud Computing Fundamentals

Projects
ResumeIQ
AI-powered resume screening and job matching platform.

Career Recommendation System
Recommendation engine based on candidate preferences.
"""


TABLE_RESUME = """Ananya Rao
Software Developer

Skills
Python | Java | SQL | Git | Docker

Education
B.Tech in Computer Science
ABC Institute of Technology

Experience
Software Developer Intern
Software Company

Projects
Employee Management System
Resume Matching Application
"""


def create_pdf(filename: str, text: str, pages: int = 1):
    """Create a synthetic PDF resume."""

    pdf_path = OUTPUT_DIR / filename

    document = pymupdf.open()

    lines = text.splitlines()

    if pages == 1:
        page = document.new_page()
        y_position = 72

        for line in lines:
            page.insert_text(
                (72, y_position),
                line,
                fontsize=11,
            )
            y_position += 16

    else:
        midpoint = len(lines) // 2

        for page_lines in (lines[:midpoint], lines[midpoint:]):
            page = document.new_page()
            y_position = 72

            for line in page_lines:
                page.insert_text(
                    (72, y_position),
                    line,
                    fontsize=11,
                )
                y_position += 16

    document.save(pdf_path)
    document.close()

    return pdf_path


def create_docx(filename: str, text: str):
    """Create a synthetic DOCX resume."""

    docx_path = OUTPUT_DIR / filename

    document = Document()

    for line in text.splitlines():
        paragraph = document.add_paragraph()

        if line.strip():
            paragraph.add_run(line)

    document.save(docx_path)

    return docx_path


def create_table_docx(filename: str):
    """Create a DOCX resume containing a table."""

    docx_path = OUTPUT_DIR / filename

    document = Document()

    document.add_heading("Ananya Rao", level=1)
    document.add_paragraph("Software Developer")

    document.add_heading("Skills", level=2)

    table = document.add_table(
        rows=2,
        cols=4,
    )

    headers = ["Python", "Java", "SQL", "Docker"]

    for index, header in enumerate(headers):
        table.cell(0, index).text = header

    values = ["Machine Learning", "Git", "REST APIs", "AWS"]

    for index, value in enumerate(values):
        table.cell(1, index).text = value

    document.add_heading("Education", level=2)
    document.add_paragraph(
        "B.Tech in Computer Science - ABC University"
    )

    document.add_heading("Experience", level=2)
    document.add_paragraph(
        "Software Developer Intern - Software Company"
    )

    document.save(docx_path)

    return docx_path


def main():
    files = []

    files.append(
        create_pdf(
            "simple_resume.pdf",
            SIMPLE_RESUME,
        )
    )

    files.append(
        create_pdf(
            "bullet_resume.pdf",
            BULLET_RESUME,
        )
    )

    files.append(
        create_pdf(
            "multi_page_resume.pdf",
            MULTI_PAGE_RESUME,
            pages=2,
        )
    )

    files.append(
        create_pdf(
            "table_style_resume.pdf",
            TABLE_RESUME,
        )
    )

    files.append(
        create_docx(
            "simple_resume.docx",
            SIMPLE_RESUME,
        )
    )

    files.append(
        create_docx(
            "bullet_resume.docx",
            BULLET_RESUME,
        )
    )

    files.append(
        create_table_docx(
            "table_resume.docx",
        )
    )

    print("\nCreated synthetic test resumes:\n")

    for file_path in files:
        print(f" - {file_path}")


if __name__ == "__main__":
    main()
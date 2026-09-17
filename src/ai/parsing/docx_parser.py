# from docx import Document
#
#
# def extract_text_from_docx(file_path: str) -> str:
#     """
#     Extract text from a DOCX resume.
#
#     Args:
#         file_path: Path to the DOCX file.
#
#     Returns:
#         Extracted paragraph text from the DOCX.
#     """
#
#     document = Document(file_path)
#
#     paragraphs = []
#
#     for paragraph in document.paragraphs:
#         text = paragraph.text.strip()
#
#         if text:
#             paragraphs.append(text)
#
#     return "\n".join(paragraphs).strip()



from docx import Document


def extract_text_from_docx(file_path: str) -> str:
    """
    Extract text from paragraphs and tables in a DOCX resume.

    Args:
        file_path: Path to the DOCX file.

    Returns:
        Combined text extracted from paragraphs and tables.
    """

    document = Document(file_path)

    text_parts = []

    # Extract paragraph text
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Extract table content
    for table in document.tables:
        for row in table.rows:
            row_text = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_text.append(cell_text)

            if row_text:
                text_parts.append(" | ".join(row_text))

    return "\n".join(text_parts).strip()